"""Unit tests for text extraction (ingestion)."""

import io

from docx import Document
from docx.oxml import parse_xml

from app.modules.ingestion.extractors import extract_segments
from app.modules.ingestion.service import _no_text_message


def _docx(build) -> bytes:
    doc = Document()
    build(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


async def test_blank_docx_yields_nothing():
    assert extract_segments(_docx(lambda d: d.add_paragraph("")), "docx") == []


async def test_docx_keeps_tables_in_their_section_and_reads_headers():
    def build(d):
        d.add_heading("Expenses", 1)
        d.add_paragraph("Submit receipts within 30 days.")
        table = d.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Grade A"
        table.cell(0, 1).text = "$500"
        d.add_heading("Travel", 1)
        d.add_paragraph("Book 14 days ahead.")
        d.sections[0].header.paragraphs[0].text = "POLICY FIN-007"

    segments = extract_segments(_docx(build), "docx")
    by_heading = {s.section_heading: s.text for s in segments}
    assert "Grade A | $500" in by_heading["Expenses"]
    assert by_heading["Travel"].endswith("Book 14 days ahead.")
    assert by_heading["Header and footer"] == "POLICY FIN-007"


async def test_docx_text_box_content_is_extracted():
    def build(d):
        run = parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            "<w:pict><wps:txbx><w:txbxContent><w:p><w:r>"
            "<w:t>Remote work is allowed two days a week.</w:t>"
            "</w:r></w:p></w:txbxContent></wps:txbx></w:pict></w:r>"
        )
        d.add_paragraph()._p.append(run)

    segments = extract_segments(_docx(build), "docx")
    assert [s.text for s in segments] == ["Remote work is allowed two days a week."]


async def test_docx_with_only_headings_is_still_indexed():
    segments = extract_segments(_docx(lambda d: d.add_heading("Onboarding", 1)), "docx")
    assert [s.text for s in segments] == ["Onboarding"]


async def test_no_text_message_depends_on_file_type():
    assert "scanned" in _no_text_message("pdf")
    assert "scanned" not in _no_text_message("docx")
